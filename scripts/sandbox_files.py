"""Générateurs de fichiers réalistes pour la sandbox et les tests (PDF, Office, images, faux fichiers créatifs)."""

from __future__ import annotations

import io
import os
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

# En-têtes binaires (« nombres magiques ») des formats que l'on ne sait pas produire réellement.
INDD_MAGIC = bytes.fromhex("0606EDF5D81D46E5BD31EFE7FE74B71D") + b"DOCUMENT"
MP4_MAGIC = b"\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00mp42isom"
OTF_MAGIC = b"OTTO\x00\x0a\x00\x80\x00\x03\x00\x20"
AEP_MAGIC = b"RIFX\x00\x00\x00\x00Egg!"


def set_mtime(path: Path, when: str | None) -> None:
    """Fixe la date de modification (format AAAA-MM-JJ)."""
    if not when:
        return
    ts = datetime.strptime(when, "%Y-%m-%d").replace(hour=10).timestamp()
    os.utime(path, (ts, ts))


def _pdf_escape(text: str) -> bytes:
    raw = text.encode("cp1252", errors="replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def write_pdf(path: Path, pages: list[list[str]]) -> None:
    """Écrit un PDF texte minimal mais valide (Helvetica, une liste de lignes par page)."""
    objects: list[bytes] = []
    n_pages = len(pages)
    page_ids = [4 + 2 * i for i in range(n_pages)]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = b" ".join(f"{pid} 0 R".encode() for pid in page_ids)
    objects.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(n_pages).encode() + b" >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, lines in enumerate(pages):
        content = b"BT /F1 11 Tf 50 790 Td 15 TL " + b" ".join(b"(" + _pdf_escape(line) + b") Tj T*" for line in lines) + b" ET"
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> /Contents "
            + str(page_ids[i] + 1).encode()
            + b" 0 R >>"
        )
        objects.append(b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(out.getvalue())


FONT_CANDIDATES = ("arial.ttf", "DejaVuSans.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
                   "/Library/Fonts/Arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")


def _font(size: int) -> ImageFont.ImageFont:
    """Police système avec les accents si possible ; sinon police intégrée de Pillow."""
    for name in FONT_CANDIDATES:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # Pillow < 10.1
        return ImageFont.load_default()


def make_image(lines: list[str], color: str = "#d9534f", size: tuple[int, int] = (800, 600)) -> Image.Image:
    """Image simple : fond coloré, quelques formes et du texte."""
    img = Image.new("RGB", size, color)
    draw = ImageDraw.Draw(img)
    w, h = size
    draw.rectangle([w * 0.08, h * 0.08, w * 0.92, h * 0.92], outline="white", width=6)
    draw.ellipse([w * 0.6, h * 0.55, w * 0.85, h * 0.85], fill="#ffffff55")
    font = _font(max(18, h // 14))
    y = h * 0.15
    for line in lines:
        draw.text((w * 0.12, y), line, fill="white", font=font)
        y += h // 9
    return img


def write_image(path: Path, lines: list[str], color: str = "#d9534f", size: tuple[int, int] = (800, 600),
                camera: str | None = None, taken: str | None = None) -> None:
    """Écrit un JPG/PNG ; ajoute des métadonnées EXIF (appareil, date de prise de vue) si fournies."""
    img = make_image(lines, color, size)
    kwargs: dict = {}
    if path.suffix.lower() in {".jpg", ".jpeg"} and (camera or taken):
        exif = Image.Exif()
        if camera:
            exif[0x0110] = camera  # Model
            exif[0x010F] = camera.split()[0]  # Make
        if taken:
            exif[0x0132] = taken.replace("-", ":") + " 10:00:00"  # DateTime
        kwargs["exif"] = exif
    img.save(path, **kwargs)


def write_scanned_pdf(path: Path, lines: list[str]) -> None:
    """PDF « scanné » : une seule image, aucun texte extractible."""
    img = Image.new("RGB", (1240, 1754), "#f4f1ea")
    draw = ImageDraw.Draw(img)
    font = _font(34)
    y = 120
    for line in lines:
        draw.text((110, y), line, fill="#222222", font=font)
        y += 62
    img.save(path, "PDF", resolution=150)


def write_psd(path: Path, lines: list[str], color: str) -> None:
    """Vrai PSD (via psd-tools) ; à défaut, faux PSD avec un en-tête valide."""
    img = make_image(lines, color, (600, 400))
    try:
        from psd_tools import PSDImage

        psd = PSDImage.frompil(img)
        with open(path, "wb") as fh:
            psd.save(fh)
    except Exception:
        path.write_bytes(b"8BPS\x00\x01" + b"\x00" * 20 + "\n".join(lines).encode("utf-8"))


def write_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(path)


def write_xlsx(path: Path, sheets: dict[str, list[list]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


def write_pptx(path: Path, slides: list[tuple[str, str]]) -> None:
    from pptx import Presentation

    prs = Presentation()
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(path)


def write_stub(path: Path, magic: bytes, text: str = "") -> None:
    """Faux fichier binaire : en-tête réaliste puis du remplissage."""
    path.write_bytes(magic + text.encode("utf-8") + b"\x00" * 2048)


def write_file(path: Path, spec: dict) -> None:
    """Aiguille vers le bon générateur selon `spec["kind"]`."""
    path.parent.mkdir(parents=True, exist_ok=True)
    kind = spec["kind"]
    lines = spec.get("lines", [])
    if kind == "pdf":
        write_pdf(path, spec.get("pages") or [lines])
    elif kind == "scan":
        write_scanned_pdf(path, lines)
    elif kind == "docx":
        write_docx(path, spec.get("title", lines[0] if lines else ""), lines[1:] if "title" not in spec else lines)
    elif kind == "xlsx":
        write_xlsx(path, spec["sheets"])
    elif kind == "pptx":
        write_pptx(path, spec["slides"])
    elif kind == "image":
        write_image(path, lines, spec.get("color", "#d9534f"), tuple(spec.get("size", (800, 600))),
                    spec.get("camera"), spec.get("taken"))
    elif kind == "psd":
        write_psd(path, lines, spec.get("color", "#2e6da4"))
    elif kind == "text":
        path.write_text(spec["text"], encoding="utf-8")
    elif kind == "copy":
        path.write_bytes(Path(spec["from"]).read_bytes())
    else:
        magic = {"indd": INDD_MAGIC, "mp4": MP4_MAGIC, "otf": OTF_MAGIC, "aep": AEP_MAGIC}[kind]
        write_stub(path, magic, " ".join(lines))
    set_mtime(path, spec.get("mtime"))
