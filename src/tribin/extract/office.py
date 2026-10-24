"""Word, Excel, PowerPoint : texte principal, feuilles et premières cellules, titres des diapositives."""

from __future__ import annotations

from pathlib import Path

from tribin.models import FileInfo

MAX_CHARS = 4000
MAX_ROWS = 25
MAX_COLS = 10


def _join(parts: list[str]) -> str:
    out, total = [], 0
    for part in parts:
        part = part.strip()
        if not part:
            continue
        out.append(part)
        total += len(part) + 1
        if total >= MAX_CHARS:
            break
    return "\n".join(out)[:MAX_CHARS]


def extract_docx(path: Path, info: FileInfo) -> None:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables[:5]:
        for row in table.rows[:MAX_ROWS]:
            parts.append(" | ".join(cell.text for cell in row.cells))
    core = doc.core_properties
    if core.title:
        info.metadata["titre"] = core.title[:200]
    if core.author:
        info.metadata["auteur"] = core.author[:100]
    info.text = _join(parts)


def extract_xlsx(path: Path, info: FileInfo) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        info.metadata["feuilles"] = ", ".join(wb.sheetnames)[:300]
        parts: list[str] = []
        for ws in wb.worksheets[:5]:
            parts.append(f"[Feuille : {ws.title}]")
            for row in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    parts.append(" | ".join(cells))
        info.text = _join(parts)
    finally:
        wb.close()


def extract_pptx(path: Path, info: FileInfo) -> None:
    from pptx import Presentation

    prs = Presentation(str(path))
    info.metadata["diapositives"] = str(len(prs.slides))
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        parts.append(f"Diapo {index} : {title}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                parts.append(shape.text_frame.text)
    info.text = _join(parts)
